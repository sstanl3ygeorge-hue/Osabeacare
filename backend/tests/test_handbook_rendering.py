"""
Regression tests for Employee Handbook rendering in agreement_document_service.py

Invariants verified without any live DB or Supabase calls:

1. No raw {{token}} placeholders survive in the rendered handbook PDF.
2. No TBC values appear in a fully-resolved handbook.
3. No draft/editorial artifact phrases survive rendering.
4. Missing required org fields raise HandbookRenderError — no broken PDF.
5. company_address maps to a real address, not the company name.
6. Rejection preserves the rendered_file_url (same handbook artifact).
"""
from __future__ import annotations

import asyncio
import io
import sys
import types
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

# ---------------------------------------------------------------------------
# Path bootstrap
# ---------------------------------------------------------------------------
BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

# ---------------------------------------------------------------------------
# Stub heavy / network-dependent imports
# ---------------------------------------------------------------------------
def _make_stub(name: str):
    mod = types.ModuleType(name)
    sys.modules.setdefault(name, mod)
    return mod


for _stub_name in [
    "motor", "motor.motor_asyncio",
    "supabase_storage",
    "services", "services.pdf_service",
]:
    _make_stub(_stub_name)

sys.modules["services.pdf_service"].get_logo_image = lambda **_kw: None  # type: ignore
_ss = sys.modules["supabase_storage"]
_ss.download_file_from_storage = MagicMock()  # type: ignore
_ss.upload_file_to_storage = MagicMock()      # type: ignore

from agreement_document_service import (  # noqa: E402
    HANDBOOK_AGREEMENT_TYPE,
    HANDBOOK_TEMPLATE_PATH,
    REQUIRED_HANDBOOK_FIELDS,
    HandbookRenderError,
    _docx_to_blocks,
    _render_pdf,
    _replace_handbook_text,
    _resolve_handbook_fields,
    _validate_handbook_fields,
    build_agreement_rendering,
    ensure_agreement_rendered,
)

# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------

_ORG_SETTINGS = {
    "organisation_name": "Osabea Healthcare Solutions Ltd",
    "organisation_address": "19 Station Road, Harlow, CM20 2BB",
    "mileage_rate": "0.45",
}

_ORG_SETTINGS_NO_ADDRESS = {
    "organisation_name": "Osabea Healthcare Solutions Ltd",
    # deliberately omitting organisation_address
}

_EMPLOYEE = {
    "id": "emp_handbook_test_001",
    "name": "Jane Smith",
    "job_title": "Care Assistant",
}

# Phrases that must NOT appear in the final rendered handbook PDF.
_FORBIDDEN_PHRASES = [
    "{{company_name}}",
    "{{company_address}}",
    "{{mileage_rate}}",
    "{{registered_manager_name}}",
    "{{grievance_contact_name}}",
    "{{grievance_contact_email}}",
    "{{about_us_text}}",
    "{{phone_number}}",
    "{{website}}",
    # Legacy template names
    "iCubeDALPro",
    "Unit 12, Harrods Road, Harlow, CM19 5BJ",
    # Draft/editorial artifacts (belt-and-suspenders check)
    "We advise that you include",
    "We advise that you give",
    "You could base these values",
    "[Add other benefits",
    "(add duration of probation)",
    "(insert Registered Manager name)",
    "(insert grievance contact name)",
    "(insert mileage rate)",
]


def _render_handbook_pdf(org_settings: dict) -> bytes:
    """Render the handbook DOCX template to PDF bytes (synchronous helper)."""
    from docx import Document

    fields = _resolve_handbook_fields(org_settings)
    doc = Document(str(HANDBOOK_TEMPLATE_PATH))
    blocks = _docx_to_blocks(doc, lambda text: _replace_handbook_text(text, fields))
    return _render_pdf(
        blocks,
        title="Employee Handbook",
        subtitle=f"{fields['company_name']} | Test",
        employee_name="Jane Smith",
    )


def _extract_pdf_text(pdf_bytes: bytes) -> str:
    from PyPDF2 import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


# ---------------------------------------------------------------------------
# Fake DB helpers (mirrors test_contract_walkthrough.py pattern)
# ---------------------------------------------------------------------------

class _FakeCollection:
    def __init__(self):
        self._docs: list[dict] = []

    async def find_one(self, query=None, projection=None):
        if not self._docs:
            return None
        # Return first doc that matches non-id query keys
        for doc in self._docs:
            if all(doc.get(k) == v for k, v in (query or {}).items() if k != "_id"):
                return {k: v for k, v in doc.items() if k != "_id"}
        return None

    async def update_one(self, query, update, upsert=False):
        for doc in self._docs:
            if all(doc.get(k) == v for k, v in query.items()):
                for k, v in update.get("$set", {}).items():
                    doc[k] = v
                return
        if upsert:
            new_doc = {**query}
            for k, v in update.get("$set", {}).items():
                new_doc[k] = v
            self._docs.append(new_doc)

    def with_doc(self, doc: dict):
        self._docs.append(doc)
        return self


class _FakeDB:
    def __init__(self, org_settings: dict):
        self.org_settings = _FakeCollection().with_doc(org_settings)
        self.contract_templates = _FakeCollection()
        self.agreement_acknowledgements = _FakeCollection()


def _stored_url_counter():
    _counter = [0]
    async def _stub(*_args, **_kwargs):
        _counter[0] += 1
        return f"https://storage.test/handbook_{_counter[0]}.pdf"
    return _stub


@contextmanager
def _patched_storage():
    with patch.multiple(
        "agreement_document_service",
        upload_file_to_storage=_stored_url_counter(),
        download_file_from_storage=AsyncMock(return_value=None),
    ):
        yield


# ---------------------------------------------------------------------------
# Test 1 — No raw placeholders in rendered handbook
# ---------------------------------------------------------------------------

class TestHandbookNoUnresolvedPlaceholders:

    def test_no_forbidden_phrases_in_rendered_pdf(self):
        """Rendered handbook must contain none of the forbidden placeholder phrases."""
        pdf_bytes = _render_handbook_pdf(_ORG_SETTINGS)
        text = _extract_pdf_text(pdf_bytes)

        for phrase in _FORBIDDEN_PHRASES:
            assert phrase not in text, (
                f"Forbidden phrase found in rendered handbook: {phrase!r}"
            )

    def test_company_name_appears_in_output(self):
        """Rendered handbook must contain the organisation name."""
        pdf_bytes = _render_handbook_pdf(_ORG_SETTINGS)
        text = _extract_pdf_text(pdf_bytes)
        assert "Osabea Healthcare Solutions Ltd" in text, (
            "Organisation name not found in rendered handbook"
        )

    def test_old_company_names_replaced(self):
        """Legacy embedded company names must not survive rendering."""
        pdf_bytes = _render_handbook_pdf(_ORG_SETTINGS)
        text = _extract_pdf_text(pdf_bytes)
        assert "iCubeDALPro" not in text, "Old company name 'iCubeDALPro' survived rendering"

    def test_mileage_rate_appears(self):
        """The mileage rate token must be replaced with the configured rate."""
        settings = {**_ORG_SETTINGS, "mileage_rate": "0.45"}
        pdf_bytes = _render_handbook_pdf(settings)
        text = _extract_pdf_text(pdf_bytes)
        assert "{{mileage_rate}}" not in text, "Mileage rate token not resolved"
        assert "0.45" in text, "Mileage rate value not found in rendered handbook"


# ---------------------------------------------------------------------------
# Test 2 — No TBC or draft artifacts
# ---------------------------------------------------------------------------

class TestHandbookQualityGuards:

    def test_no_tbc_in_fully_resolved_handbook(self):
        """A handbook with complete org settings must contain no 'TBC'."""
        pdf_bytes = _render_handbook_pdf(_ORG_SETTINGS)
        text = _extract_pdf_text(pdf_bytes)
        assert "TBC" not in text, "TBC value found in fully-resolved handbook PDF"

    def test_no_editorial_advisory_text(self):
        """Draft/editorial instructions must not appear in the rendered PDF."""
        pdf_bytes = _render_handbook_pdf(_ORG_SETTINGS)
        text = _extract_pdf_text(pdf_bytes)
        advisory_phrases = [
            "We advise that you include",
            "We advise that you give",
            "You could base these values",
            "[Add other benefits",
            "(add duration of probation)",
        ]
        for phrase in advisory_phrases:
            assert phrase not in text, (
                f"Draft/editorial artifact phrase found in rendered handbook: {phrase!r}"
            )

    def test_company_address_is_address_not_company_name(self):
        """company_address field must be a street address, not the company name."""
        fields = _resolve_handbook_fields(_ORG_SETTINGS)
        addr = fields.get("company_address", "")
        name = fields.get("company_name", "")
        assert addr != name, f"company_address equals company_name: {addr!r}"
        assert "Osabea Healthcare Solutions Ltd" not in addr, (
            f"Company name leaked into company_address: {addr!r}"
        )
        assert "Station Road" in addr, (
            f"Expected street address in company_address, got: {addr!r}"
        )


# ---------------------------------------------------------------------------
# Test 3 — Missing required fields raise HandbookRenderError
# ---------------------------------------------------------------------------

class TestHandbookFieldValidation:

    def test_complete_org_settings_no_error(self):
        """Complete org settings must not raise."""
        fields = _resolve_handbook_fields(_ORG_SETTINGS)
        _validate_handbook_fields(fields)  # must not raise

    def test_missing_company_address_raises(self):
        """Org settings without an address must raise HandbookRenderError."""
        fields = _resolve_handbook_fields(_ORG_SETTINGS_NO_ADDRESS)
        with pytest.raises(HandbookRenderError, match="company_address"):
            _validate_handbook_fields(fields)

    def test_missing_company_name_raises(self):
        """Org settings without a company name must raise HandbookRenderError."""
        fields = _resolve_handbook_fields({})  # empty settings
        with pytest.raises(HandbookRenderError, match="company_name"):
            _validate_handbook_fields(fields)

    @pytest.mark.parametrize("missing_field", REQUIRED_HANDBOOK_FIELDS)
    def test_each_required_field_independently(self, missing_field):
        """Every REQUIRED_HANDBOOK_FIELDS entry must individually raise when blanked."""
        fields = _resolve_handbook_fields(_ORG_SETTINGS)
        fields[missing_field] = ""  # forcibly blank
        with pytest.raises(HandbookRenderError, match=missing_field):
            _validate_handbook_fields(fields)

    def test_missing_address_blocks_full_render(self):
        """end-to-end: missing org address must raise HandbookRenderError before producing a PDF."""
        with _patched_storage():
            db = _FakeDB(_ORG_SETTINGS_NO_ADDRESS)
            with pytest.raises(HandbookRenderError):
                asyncio.run(build_agreement_rendering(db, _EMPLOYEE, HANDBOOK_AGREEMENT_TYPE))


# ---------------------------------------------------------------------------
# Test 4 — Worker/admin use the same rendered handbook artifact
# ---------------------------------------------------------------------------

class TestHandbookArtifactConsistency:

    def test_same_url_across_calls_same_version(self):
        """Second call with same template version reuses the existing rendered URL."""
        with _patched_storage():
            db = _FakeDB(_ORG_SETTINGS)

            result1 = asyncio.run(
                ensure_agreement_rendered(db, _EMPLOYEE, HANDBOOK_AGREEMENT_TYPE)
            )
            url1 = result1.get("rendered_file_url")
            assert url1, "First render did not produce a rendered_file_url"

            result2 = asyncio.run(
                ensure_agreement_rendered(db, _EMPLOYEE, HANDBOOK_AGREEMENT_TYPE)
            )
            url2 = result2.get("rendered_file_url")
            assert url1 == url2, (
                f"Second call produced a different URL: {url1!r} vs {url2!r}. "
                "Worker and admin must reference the same handbook artifact."
            )

    def test_rejection_preserves_rendered_pdf_url(self):
        """After rejection the rendered_file_url must still be present."""
        with _patched_storage():
            db = _FakeDB(_ORG_SETTINGS)

            result = asyncio.run(
                ensure_agreement_rendered(db, _EMPLOYEE, HANDBOOK_AGREEMENT_TYPE)
            )
            original_url = result.get("rendered_file_url")
            assert original_url

            # Simulate admin rejection by setting verification_status
            asyncio.run(
                db.agreement_acknowledgements.update_one(
                    {"employee_id": _EMPLOYEE["id"], "agreement_type": HANDBOOK_AGREEMENT_TYPE},
                    {"$set": {"verification_status": "rejected", "acknowledged": False}},
                )
            )

            # Re-render (same template version) — must reuse existing URL
            result2 = asyncio.run(
                ensure_agreement_rendered(db, _EMPLOYEE, HANDBOOK_AGREEMENT_TYPE)
            )
            assert result2.get("rendered_file_url") == original_url, (
                "Rejection caused the rendered_file_url to change — worker would lose PDF access"
            )
